# -*- coding: utf-8 -*-
"""Patch Alex's manually edited deck: A -> A1/A2 on the models slide, new A2 example slide."""
import shutil
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

IN = 914400
SLIDE_W = 24387175
X0 = 1955800
CW = SLIDE_W - 2 * X0
TITLE_Y = 1133475
RED = 'C00000'
INK = '101010'
GREY = '595959'
LINEGREY = 'E3E3E3'
PANEL = 'F4F4F4'
WHITE = 'FFFFFF'
FB = 'e-Ukraine Bold'
FL = 'e-Ukraine Light'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'

def E(v):
    return int(v * IN)

def R(t, sz=2000, font=FL, color=INK, bold=False):
    return dict(t=t, sz=sz, font=font, color=color, bold=bold)

def P(runs, align=None, spc_bef=None, ln_pts=None, bullet=False):
    if isinstance(runs, dict):
        runs = [runs]
    return dict(runs=runs, align=align, spc_bef=spc_bef, ln_pts=ln_pts, bullet=bullet)

def make_p(spec):
    p = etree.Element(qn('a:p'))
    pPr = etree.SubElement(p, qn('a:pPr'))
    if spec['align']:
        pPr.set('algn', spec['align'])
    if spec['ln_pts']:
        ln = etree.SubElement(pPr, qn('a:lnSpc'))
        etree.SubElement(ln, qn('a:spcPts')).set('val', str(spec['ln_pts']))
    if spec['spc_bef']:
        sb = etree.SubElement(pPr, qn('a:spcBef'))
        etree.SubElement(sb, qn('a:spcPts')).set('val', str(spec['spc_bef']))
    if spec['bullet']:
        pPr.set('marL', '342900')
        pPr.set('indent', '-342900')
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
    else:
        etree.SubElement(pPr, qn('a:buNone'))
    for r in spec['runs']:
        rEl = etree.SubElement(p, qn('a:r'))
        rPr = etree.SubElement(rEl, qn('a:rPr'))
        rPr.set('lang', 'uk-UA')
        rPr.set('sz', str(r['sz']))
        rPr.set('dirty', '0')
        if r['bold']:
            rPr.set('b', '1')
        fill = etree.SubElement(rPr, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr')).set('val', r['color'])
        etree.SubElement(rPr, qn('a:latin')).set('typeface', r['font'])
        tEl = etree.SubElement(rEl, qn('a:t'))
        tEl.text = r['t']
        if r['t'] != r['t'].strip():
            tEl.set(XML_SPACE, 'preserve')
    return p

def set_paras(owner, paras, anchor='t', wrap=True):
    tf = owner.text_frame
    txBody = tf._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    bodyPr = txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bodyPr.set(k, '0')
    bodyPr.set('wrap', 'square' if wrap else 'none')
    bodyPr.set('anchor', {'t': 't', 'm': 'ctr', 'b': 'b'}[anchor])
    for tag in ('a:normAutofit', 'a:spAutoFit'):
        f = bodyPr.find(qn(tag))
        if f is not None:
            bodyPr.remove(f)
    for spec in paras:
        txBody.append(make_p(spec))

def add_box(slide, x, y, w, h, paras, anchor='t'):
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    set_paras(tb, paras, anchor=anchor)
    return tb

def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)

def add_card(slide, x, y, w, h, fill=WHITE, line=LINEGREY):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    sh.adjustments[0] = 0.06
    if fill is None:
        sh.fill.background()
    else:
        _fill(sh, fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Emu(12700)
    sh.shadow.inherit = False
    set_paras(sh, [P(R('', sz=200))])
    return sh

def add_table(slide, rows, cols, x, y, w, h):
    gf = slide.shapes.add_table(rows, cols, Emu(int(x)), Emu(int(y)),
                                Emu(int(w)), Emu(int(h)))
    tbl = gf.table
    tblPr = tbl._tbl.find(qn('a:tblPr'))
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    return tbl

def cell_set(cell, paras, fill=WHITE, anchor='m'):
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE}[anchor]
    cell.margin_left = Emu(E(0.10))
    cell.margin_right = Emu(E(0.10))
    cell.margin_top = Emu(E(0.045))
    cell.margin_bottom = Emu(E(0.045))
    txBody = cell._tc.find(qn('a:txBody'))
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    for spec in paras:
        txBody.append(make_p(spec))

def cell_borders(cell, color=LINEGREY, w_emu=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    order = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']
    for tag in order:
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for tag in order:
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set('w', str(w_emu))
        ln.set('cap', 'flat')
        f = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(f, qn('a:srgbClr')).set('val', color)
    lns = [c for c in tcPr if c.tag in {qn(t) for t in order}]
    for ln in reversed(lns):
        tcPr.remove(ln)
        tcPr.insert(0, ln)

def iter_all_runs(slide):
    """Yield (run_element) across shapes and table cells."""
    def walk_txbody(txBody):
        for p in txBody.findall(qn('a:p')):
            for r in p.findall(qn('a:r')):
                yield r
    for shp in slide.shapes:
        if shp.has_text_frame:
            yield from walk_txbody(shp.text_frame._txBody)
        if shp.shape_type == 19 or (shp.element.tag.endswith('graphicFrame')):
            for tc in shp.element.iter(qn('a:tc')):
                tx = tc.find(qn('a:txBody'))
                if tx is not None:
                    yield from walk_txbody(tx)

def replace_run_text(slide, old, new, count=1):
    n = 0
    for r in iter_all_runs(slide):
        t = r.find(qn('a:t'))
        if t is not None and t.text and old in t.text:
            t.text = t.text.replace(old, new)
            n += 1
            if n >= count:
                return n
    return n

def slide_text(slide):
    return ' | '.join(t.text or '' for r in iter_all_runs(slide)
                      for t in [r.find(qn('a:t'))] if t is not None)

# ---------------- open Alex's current file ----------------
shutil.copy('current.pptx', 'patched.pptx')
prs = Presentation('patched.pptx')

slides = list(prs.slides)
s_models = next(s for s in slides if 'моделі атрибуції витрат' in slide_text(s))
s_a1 = next(s for s in slides if 'Модель A1' in slide_text(s))
idx_a1 = slides.index(s_a1)

# ---------------- 1. Models slide: A column -> A1 + A2 ----------------
# delete the old A card + its text boxes (keep panel + header band)
for shp in list(s_models.shapes):
    l, t = shp.left / IN, shp.top / IN
    r_edge = (shp.left + (shp.width or 0)) / IN
    if 2.25 <= l and r_edge <= 7.45 and t >= 4.55:
        shp._element.getparent().remove(shp._element)

CX = 2.3888
CARD_W = 4.9
Y1, H1 = 4.69, 3.32
Y2 = Y1 + H1 + 0.31
twins = [
    (Y1, 'A1', 'Пряма атрибуція FTE',
     'ФОТ Delivery і Support',
     'FTE-частки × ЗП команд = собівартість компонента'),
    (Y2, 'A2', 'Закупівлі у постачальника',
     'Cubbit S3: постачальник тарифікує кожен GB',
     'Units included × ціна юніта = собівартість компонента'),
]
for cy, let, name, ex, mech in twins:
    add_card(s_models, E(CX), E(cy), E(CARD_W), E(H1))
    bx, bw = CX + 0.26, CARD_W - 0.52
    add_box(s_models, E(bx), E(cy + 0.18), E(bw), E(0.52),
            [P(R(let, sz=2800, font=FB, color=RED), ln_pts=2900)])
    add_box(s_models, E(bx), E(cy + 0.72), E(bw), E(0.62),
            [P(R(name, sz=1900, font=FB, color=INK), ln_pts=2200)])
    add_box(s_models, E(bx), E(cy + 1.36), E(bw), E(0.62),
            [P(R(ex, sz=1600, font=FL, color=GREY), ln_pts=1950)])
    add_box(s_models, E(bx), E(cy + 2.02), E(bw), E(0.32),
            [P(R('У ціну компонента', sz=1600, font=FB, color=RED))])
    add_box(s_models, E(bx), E(cy + 2.36), E(bw), E(0.9),
            [P(R(mech, sz=1700, font=FL, color=INK), ln_pts=2050)])

# retitle the models slide (count-neutral)
replace_run_text(s_models, '4 моделі атрибуції витрат', 'Моделі атрибуції витрат')

# ---------------- 2. New slide: Модель A2 — приклад ----------------
layout1 = None
for lay in prs.slide_masters[0].slide_layouts:
    if lay.part.partname.endswith('slideLayout1.xml'):
        layout1 = lay
        break
s = prs.slides.add_slide(layout1)

add_box(s, X0, TITLE_Y, CW, E(1.75),
        [P(R('Модель A2 — приклад', sz=9600, font=FB, color=INK), ln_pts=9600)])
add_box(s, X0, E(3.05), CW, E(0.8),
        [P(R('Прямі закупівлі в собівартість: скільки юнітів постачальника входить в один проданий компонент.',
             sz=2100, font=FL, color=GREY), ln_pts=2500)])

COLS = ['Компонент', 'Юніт постачальника', 'Units included\n(юнітів на 1 компонент)',
        'Ціна юніта, ₴/міс', 'COGS = units × ціна', 'Ціна компонента', 'Маржа']
W = [4.3, 3.5, 3.6, 2.5, 3.4, 2.6, 2.49]
t = add_table(s, 3, 7, X0, E(4.3), CW, E(2.6))
for i, wdt in enumerate(W):
    t.columns[i].width = Emu(E(wdt))
hdr = ['Компонент', 'Юніт постачальника', 'Units included, на 1 компонент',
       'Ціна юніта, ₴/міс', 'COGS = units × ціна', 'Ціна компонента', 'Маржа']
for c, name in enumerate(hdr):
    cell_set(t.cell(0, c), [P(R(name, sz=1600, font=FB, color=WHITE),
                              align='l' if c == 0 else 'ctr', ln_pts=1850)], fill=RED)
row = ['Cubbit S3 Storage, 10 TB', '1 GB — погігабайтна тарифікація', '10 000',
       '0,35', '3 500 ₴ ≈ €72', '€85,00', '≈ €13 · 15 %']
for c, v in enumerate(row):
    bold = c in (4, 6)
    cell_set(t.cell(1, c), [P(R(v, sz=1800, font=FB if bold else FL, color=INK),
                              align='l' if c == 0 else 'ctr', ln_pts=2100)], fill=WHITE)
for c in range(7):
    cell_set(t.cell(2, c), [P(R('…', sz=1800, font=FL, color=GREY),
                              align='l' if c == 0 else 'ctr', ln_pts=2100)], fill=WHITE)
for ri in range(3):
    for ci in range(7):
        cell_borders(t.cell(ri, ci))

add_box(s, X0, E(8.3), CW - E(3.5), E(3.2), [
    P([R('Механіка: ', sz=2000, font=FB, color=INK),
       R('компонент = 10 TB = 10 000 GB, постачальник тарифікує кожен GB → units included = 10 000. '
         'COGS = 10 000 × 0,35 = 3 500 ₴/міс.', sz=2000, font=FL, color=INK)], ln_pts=2550),
    P([R('Перевірка: ', sz=2000, font=FB, color=INK),
       R('Σ (units × ціна юніта) по проданих компонентах ≈ рахунки постачальника за період.',
         sz=2000, font=FL, color=INK)], spc_bef=550, ln_pts=2550),
])
bar = add_card(s, X0, E(12.8), E(18.9), E(1.5), fill=PANEL, line=None)
set_paras(bar, [P([
    R('Джерело: ', sz=1700, font=FB, color=RED),
    R('CRM — Cubbit S3 Storage, 10 TB (у CRM поле зберігає обернене співвідношення 0,0001 компонента '
      'на 1 GB; у розрахунку — 10 000 юнітів на компонент). Курс 48,5 ₴/€ — ілюстративний; '
      'правило приведення валют фіксується з фінансами.', sz=1700, font=FL, color=INK)],
    ln_pts=1950)], anchor='m')
bp = bar.text_frame._txBody.find(qn('a:bodyPr'))
bp.set('lIns', str(E(0.35))); bp.set('rIns', str(E(0.3)))

# move the new slide right after A1 example
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_el = ids[-1]
sldIdLst.remove(new_el)
sldIdLst.insert(idx_a1 + 1, new_el)

# ---------------- 3. Cross-references ----------------
s_stack = next(x for x in prs.slides if 'Складові ціни компонента' in slide_text(x))
replace_run_text(s_stack,
                 'обладнання, ліцензії, колокейшн, канали + ФОТ доставки і супроводу (модель A)',
                 'закупівлі — обладнання, ліцензії, колокейшн, канали (A2) + ФОТ доставки і супроводу (A1)')
s_steps = next(x for x in prs.slides if 'Наступні кроки' in slide_text(x))
replace_run_text(s_steps, 'модель A / B / C / D', 'модель A1 / A2 / B / C / D')

prs.save('patched.pptx')
print('patched; slides:', len(prs.slides._sldIdLst))
