# -*- coding: utf-8 -*-
"""Patch v4: component-bridge base -> average COGS weight in average product configuration.
Identical note on models slide + B/C/D example slides; purge price-share wording."""
import shutil
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from lxml import etree

IN = 914400
RED = 'C00000'
INK = '101010'
FB = 'e-Ukraine Bold'
FL = 'e-Ukraine Light'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'

def E(v):
    return int(v * IN)

def R(t, sz=1700, font=FL, color=INK, bold=False):
    return dict(t=t, sz=sz, font=font, color=color, bold=bold)

def P(runs, align=None, spc_bef=None, ln_pts=None):
    if isinstance(runs, dict):
        runs = [runs]
    return dict(runs=runs, align=align, spc_bef=spc_bef, ln_pts=ln_pts)

def make_p(spec):
    p = etree.Element(qn('a:p'))
    pPr = etree.SubElement(p, qn('a:pPr'))
    if spec['align']:
        pPr.set('algn', spec['align'])
    if spec['ln_pts']:
        ln = etree.SubElement(pPr, qn('a:lnSpc'))
        etree.SubElement(ln, qn('a:spcPts')).set('val', str(spec['ln_pts']))
    if spec['spc_bef']:
        sb = etree.SubElement(pPr, qn('a:spcBef'))
        etree.SubElement(sb, qn('a:spcPts')).set('val', str(spec['spc_bef']))
    etree.SubElement(pPr, qn('a:buNone'))
    for r in spec['runs']:
        rEl = etree.SubElement(p, qn('a:r'))
        rPr = etree.SubElement(rEl, qn('a:rPr'))
        rPr.set('lang', 'uk-UA')
        rPr.set('sz', str(r['sz']))
        rPr.set('dirty', '0')
        if r['bold']:
            rPr.set('b', '1')
        fill = etree.SubElement(rPr, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr')).set('val', r['color'])
        etree.SubElement(rPr, qn('a:latin')).set('typeface', r['font'])
        tEl = etree.SubElement(rEl, qn('a:t'))
        tEl.text = r['t']
        if r['t'] != r['t'].strip():
            tEl.set(XML_SPACE, 'preserve')
    return p

def set_paras(shape, paras, anchor='m'):
    txBody = shape.text_frame._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    bodyPr = txBody.find(qn('a:bodyPr'))
    bodyPr.set('anchor', {'t': 't', 'm': 'ctr'}[anchor])
    for spec in paras:
        txBody.append(make_p(spec))

def find_shape(slide, marker):
    for shp in slide.shapes:
        if shp.has_text_frame and marker in shp.text_frame.text:
            return shp
    return None

def slide_with(prs, marker):
    for s in prs.slides:
        for shp in s.shapes:
            if shp.has_text_frame and marker in shp.text_frame.text:
                return s
    return None

NOTE_LEAD = 'До рівня компонента: '
NOTE_BODY = ('атрибуцію рівня продукту чи групи розкладаємо на компоненти за середньою вагою '
             'COGS компонента в середній конфігурації продукту. Окремо проаналізуємо, чи рахувати '
             'середні конфігурації за сегментами клієнтів окремо (Enterprise vs менший клієнт).')

def note_para(sz=1700, spc_bef=None):
    return P([R(NOTE_LEAD, sz=sz, font=FB, color=RED),
              R(NOTE_BODY, sz=sz, font=FL, color=INK)], ln_pts=int(sz * 1.18), spc_bef=spc_bef)

shutil.copy('current2.pptx', 'patched2.pptx')
prs = Presentation('patched2.pptx')

# ---- 1. Models slide: bottom bar -> the standard note; fix B/C/D bridge rows ----
s_models = slide_with(prs, 'Моделі атрибуції витрат')
bar = find_shape(s_models, 'Міст до компонента')
set_paras(bar, [note_para(sz=1700)], anchor='m')

bridges = {
    'Частка компонента в ціні продукту': 'CAC продукту → компоненти за середньою вагою COGS',
    'Пул ÷ дохід групи': 'Пул групи → компоненти за середньою вагою COGS',
    'Погоджена база (тут — % маржі груп)': 'Погоджена база (тут — % маржі груп) → далі за вагою COGS',
}
for old, new in bridges.items():
    shp = find_shape(s_models, old)
    if shp is None:
        print('!! not found on models slide:', old)
        continue
    set_paras(shp, [P(R(new, sz=1800, font=FL, color=INK), ln_pts=2150)], anchor='t')

# ---- 2. Model B example: caption -> mechanics (no price-share) + note ----
s_b = slide_with(prs, 'Модель B — приклад')
cap = find_shape(s_b, 'Механіка: ')
set_paras(cap, [
    P([R('Механіка: ', sz=1700, font=FB, color=RED),
       R('CAC = ФОТ каналу, віднесений на продажі продукту; «CAC з LT» = CAC ÷ LT — '
         'щомісячна складова ціни. Цифри ілюстративні (аркуш Models).',
         sz=1700, font=FL, color=INK)], ln_pts=2000),
    note_para(sz=1700, spc_bef=350),
], anchor='m')
cap.top = Emu(E(12.55))
cap.height = Emu(E(1.75))

# ---- 3. Model C example: caption -> drop price-per-component claim + note ----
s_c = slide_with(prs, 'Модель C — приклад')
cap = find_shape(s_c, 'Приклад: ')
set_paras(cap, [
    P([R('Приклад: ', sz=1700, font=FB, color=RED),
       R('ЗП ВОК — факт Q2’26 (1 030 714 ₴/кв → 343 571 ₴/міс), доходи груп — аркуш «Продукти». '
         'Частку ВОК на Resell для ілюстрації взято 100 % — реальну дасть FTE-опитування.',
         sz=1700, font=FL, color=INK)], ln_pts=2000),
    note_para(sz=1700, spc_bef=350),
], anchor='m')
cap.top = Emu(E(12.55))
cap.height = Emu(E(1.75))

# ---- 4. Model D example: caption + note ----
s_d = slide_with(prs, 'Модель D — приклад')
cap = find_shape(s_d, 'Джерело: ')
set_paras(cap, [
    P([R('Джерело: ', sz=1700, font=FB, color=RED),
       R('аркуш Models — прогноз FY2026 (маржа 50 млн у.о.), кількість продажів з прикладу '
         'моделі B. Цифри ілюстративні.', sz=1700, font=FL, color=INK)], ln_pts=2000),
    note_para(sz=1700, spc_bef=350),
], anchor='m')
cap.top = Emu(E(12.55))
cap.height = Emu(E(1.75))

prs.save('patched2.pptx')
print('patched2 saved')
