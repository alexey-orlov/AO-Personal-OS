# -*- coding: utf-8 -*-
"""Patch v5: output-table slide after 'Мета і ключові результати' + step 10 (min profitability)."""
import copy as cpmod
import shutil
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

IN = 914400
SLIDE_W = 24387175
X0 = 1955800
CW = SLIDE_W - 2 * X0
TITLE_Y = 1133475
RED = 'C00000'
INK = '101010'
GREY = '595959'
LTGREY = 'BFBFBF'
LINEGREY = 'E3E3E3'
PANEL = 'F4F4F4'
WHITE = 'FFFFFF'
FB = 'e-Ukraine Bold'
FL = 'e-Ukraine Light'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'

def E(v):
    return int(v * IN)

def R(t, sz=2000, font=FL, color=INK, bold=False):
    return dict(t=t, sz=sz, font=font, color=color, bold=bold)

def P(runs, align=None, spc_bef=None, ln_pts=None, bullet=False):
    if isinstance(runs, dict):
        runs = [runs]
    return dict(runs=runs, align=align, spc_bef=spc_bef, ln_pts=ln_pts, bullet=bullet)

def make_p(spec):
    p = etree.Element(qn('a:p'))
    pPr = etree.SubElement(p, qn('a:pPr'))
    if spec['align']:
        pPr.set('algn', spec['align'])
    if spec['ln_pts']:
        ln = etree.SubElement(pPr, qn('a:lnSpc'))
        etree.SubElement(ln, qn('a:spcPts')).set('val', str(spec['ln_pts']))
    if spec['spc_bef']:
        sb = etree.SubElement(pPr, qn('a:spcBef'))
        etree.SubElement(sb, qn('a:spcPts')).set('val', str(spec['spc_bef']))
    if spec['bullet']:
        pPr.set('marL', '342900')
        pPr.set('indent', '-342900')
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
    else:
        etree.SubElement(pPr, qn('a:buNone'))
    for r in spec['runs']:
        rEl = etree.SubElement(p, qn('a:r'))
        rPr = etree.SubElement(rEl, qn('a:rPr'))
        rPr.set('lang', 'uk-UA')
        rPr.set('sz', str(r['sz']))
        rPr.set('dirty', '0')
        if r['bold']:
            rPr.set('b', '1')
        fill = etree.SubElement(rPr, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr')).set('val', r['color'])
        etree.SubElement(rPr, qn('a:latin')).set('typeface', r['font'])
        tEl = etree.SubElement(rEl, qn('a:t'))
        tEl.text = r['t']
        if r['t'] != r['t'].strip():
            tEl.set(XML_SPACE, 'preserve')
    return p

def set_paras(owner, paras, anchor='t', wrap=True):
    txBody = owner.text_frame._txBody
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    bodyPr = txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bodyPr.set(k, '0')
    bodyPr.set('wrap', 'square' if wrap else 'none')
    bodyPr.set('anchor', {'t': 't', 'm': 'ctr'}[anchor])
    for tag in ('a:normAutofit', 'a:spAutoFit'):
        f = bodyPr.find(qn(tag))
        if f is not None:
            bodyPr.remove(f)
    for spec in paras:
        txBody.append(make_p(spec))

def add_box(slide, x, y, w, h, paras, anchor='t'):
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    set_paras(tb, paras, anchor=anchor)
    return tb

def add_rect(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    set_paras(sh, [P(R('', sz=200))])
    return sh

def add_table(slide, rows, cols, x, y, w, h):
    gf = slide.shapes.add_table(rows, cols, Emu(int(x)), Emu(int(y)),
                                Emu(int(w)), Emu(int(h)))
    tbl = gf.table
    tblPr = tbl._tbl.find(qn('a:tblPr'))
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    return tbl

def cell_set(cell, paras, fill=WHITE, mL=0.08, mR=0.08):
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Emu(E(mL))
    cell.margin_right = Emu(E(mR))
    cell.margin_top = Emu(E(0.04))
    cell.margin_bottom = Emu(E(0.04))
    txBody = cell._tc.find(qn('a:txBody'))
    for p in txBody.findall(qn('a:p')):
        txBody.remove(p)
    for spec in paras:
        txBody.append(make_p(spec))

def cell_borders(cell, color=LINEGREY, w_emu=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    order = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']
    for tag in order:
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for tag in order:
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set('w', str(w_emu))
        ln.set('cap', 'flat')
        f = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(f, qn('a:srgbClr')).set('val', color)
    lns = [c for c in tcPr if c.tag in {qn(t) for t in order}]
    for ln in reversed(lns):
        tcPr.remove(ln)
        tcPr.insert(0, ln)

def slide_with(prs, marker):
    for s in prs.slides:
        for shp in s.shapes:
            if shp.has_text_frame and marker in shp.text_frame.text:
                return s
    return None

shutil.copy('current3.pptx', 'patched3.pptx')
prs = Presentation('patched3.pptx')

# ---------------- 1. New output-table slide after 'Мета і ключові результати' ----------------
layout1 = None
for lay in prs.slide_masters[0].slide_layouts:
    if lay.part.partname.endswith('slideLayout1.xml'):
        layout1 = lay
        break
s = prs.slides.add_slide(layout1)

add_box(s, X0, TITLE_Y, CW, E(1.75),
        [P(R('Мета і ключові результати', sz=9600, font=FB, color=INK), ln_pts=9600)])
add_box(s, X0, E(3.0), CW, E(0.7),
        [P(R('Вихідний артефакт проєкту — розрахункова таблиця по кожному компоненту:',
             sz=2100, font=FL, color=GREY), ln_pts=2500)])

# column model: (header, width_in, type)  types: comp/cogs/fixed/decision/calc
COLS = [
    ('Компонент', 3.4, 'comp'),
    ('COGS, грн', 1.75, 'cogs'),
    ('CAC, грн', 1.75, 'calc'),
    ('Загальні, грн', 1.75, 'calc'),
    ('Поточна ціна, грн', 1.9, 'fixed'),
    ('Прибуток зараз, грн', 1.9, 'calc'),
    ('Прибуток зараз, %', 1.75, 'calc'),
    ('Мін. прибутко-вість, %', 2.0, 'decision'),
    ('Нова ціна, грн', 1.85, 'calc'),
    ('Зміна ціни, %', 1.7, 'calc'),
    ('Мін. ціна продажу, грн', 2.64, 'calc'),
]
TYPE_FILL = {'comp': (PANEL, INK), 'cogs': (GREY, WHITE), 'fixed': (LTGREY, INK),
             'decision': (INK, WHITE), 'calc': (RED, WHITE)}
rows_data = [
    ['Node for Private Cloud on VMware', '30 000', '4 000', '6 000', '45 000',
     '5 000', '11 %', '10 %', '50 000', '+11 %', '44 444'],
    ['vCPU (Public Cloud), 1 GHz', '60', '10', '14', '180',
     '96', '53 %', '10 %', '120', '−33 %', '93'],
    ['…'] * 11,
]
t = add_table(s, 4, 11, X0, E(3.9), CW, E(3.0))
for i, (name, wdt, typ) in enumerate(COLS):
    t.columns[i].width = Emu(E(wdt))
    fill, tcol = TYPE_FILL[typ]
    t.cell(0, i).text_frame  # ensure
    cell_set(t.cell(0, i), [P(R(name.replace('-', ''), sz=1400, font=FB, color=tcol),
                              align='l' if i == 0 else 'ctr', ln_pts=1650)], fill=fill)
t.rows[0].height = Emu(E(0.95))
for ri, row in enumerate(rows_data, start=1):
    t.rows[ri].height = Emu(E(0.68))
    for c, v in enumerate(row):
        dots = v == '…'
        bold = c in (8, 9) and not dots
        cell_set(t.cell(ri, c),
                 [P(R(v, sz=1600, font=FB if bold else FL,
                      color=INK if not dots else GREY),
                    align='l' if c == 0 else 'ctr', ln_pts=1900)], fill=WHITE)
for ri in range(4):
    for ci in range(11):
        cell_borders(t.cell(ri, ci))

# legend chips
legend = [
    (RED, 'Обчислюємо за моделями A1–D'),
    (GREY, 'Вже пораховано — перевіримо повторно'),
    (LTGREY, 'Фіксовано на час проєкту'),
    (INK, 'Політичне рішення проєкту'),
]
slot = CW / 4
for i, (fill, label) in enumerate(legend):
    cx = X0 + int(i * slot)
    add_rect(s, cx, E(7.62), E(0.4), E(0.4), fill)
    add_box(s, cx + E(0.58), E(7.58), int(slot) - E(0.7), E(0.75),
            [P(R(label, sz=1500, font=FL, color=INK), ln_pts=1800)])

# footnotes
add_box(s, X0, E(8.85), CW, E(0.5),
        [P(R('Як читати', sz=1800, font=FB, color=RED))])
notes = [
    'Прибуток зараз = поточна ціна − COGS − CAC − загальні витрати (грн і % від поточної ціни).',
    'CAC і загальні — за моделями атрибуції; one-time частини приведені до місяця через customer lifetime.',
    'Мін. прибутковість — % від фактичної ціни продажу (не прайсової: можливі знижки); погоджують CFO, CEO, CBDO.',
    'Мін. ціна продажу = (COGS + CAC + загальні) ÷ (1 − мін. прибутковість) — нижня межа для будь-якої знижки.',
    'Нова ціна — розрахунок за цільовою прибутковістю; зміна ціни — відносно поточної. Цифри в таблиці ілюстративні.',
]
paras = []
for n in notes:
    paras.append(P(R(n, sz=1700, font=FL, color=INK), bullet=True,
                   spc_bef=300, ln_pts=2050))
add_box(s, X0, E(9.45), CW, E(3.4), paras)

# position: right after 'Мета і ключові результати' (index 1) -> insert at 2
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_el = ids[-1]
sldIdLst.remove(new_el)
sldIdLst.insert(2, new_el)

# ---------------- 2. Steps slide: add step 10 ----------------
s_steps = slide_with(prs, 'Наступні кроки')
tblshape = None
for shp in s_steps.shapes:
    if shp.has_table:
        tblshape = shp
        break
tbl = tblshape.table
trs = tbl._tbl.findall(qn('a:tr'))
new_tr = cpmod.deepcopy(trs[-1])
tbl._tbl.append(new_tr)
nrow = len(tbl.rows) - 1
cell_set(tbl.cell(nrow, 0), [P(R('10', sz=1900, font=FB, color=INK), align='ctr', ln_pts=2200)], fill=WHITE, mL=0.10, mR=0.10)
cell_set(tbl.cell(nrow, 1), [P(R('Мінімальна прибутковість', sz=1800, font=FB, color=INK), align='l', ln_pts=2100)], fill=WHITE, mL=0.10, mR=0.10)
cell_set(tbl.cell(nrow, 2), [P(R('Закладаємо мінімально допустиму прибутковість — «в нуль» не продаємо; погоджують CFO, CEO, CBDO',
                                 sz=1700, font=FL, color=INK), align='l', ln_pts=2000)], fill=WHITE, mL=0.10, mR=0.10)
for ci in range(3):
    cell_borders(tbl.cell(nrow, ci))
# compact rows so the longer table still fits
tblshape.top = Emu(E(3.3))
for ri in range(len(tbl.rows)):
    h = tbl.rows[ri].height / IN
    if h > 1.2:
        tbl.rows[ri].height = Emu(E(1.32))
    elif h > 0.6:
        tbl.rows[ri].height = Emu(E(0.66))
    else:
        tbl.rows[ri].height = Emu(E(0.5))

prs.save('patched3.pptx')
print('patched3 saved; slides:', len(prs.slides._sldIdLst))
