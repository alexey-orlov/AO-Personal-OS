import sys
from pptx import Presentation
IN=914400
p = Presentation(sys.argv[1]); out = []
for i, s in enumerate(p.slides, 1):
    out.append(f"\n==================== SLIDE {i} (layout: {s.slide_layout.name})")
    items = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            items.append((sh.top, sh.left, 'T', sh.text_frame.text.strip().replace('\n', ' / ')))
        if sh.has_table:
            rows = []
            for r in sh.table.rows:
                rows.append(' | '.join(c.text.replace('\n', ' / ') for c in r.cells))
            items.append((sh.top, sh.left, 'TABLE', '\n      ' + '\n      '.join(rows)))
    for top, left, kind, txt in sorted(items):
        out.append(f"  [{kind} y={top/IN:.2f} x={left/IN:.2f}] {txt}")
    if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
        out.append(f"  NOTES: {s.notes_slide.notes_text_frame.text}")
open(sys.argv[2], 'w').write('\n'.join(out)); print('ok', len(p.slides))
