"""Oracle packs · packaging strategy deck (SoftServe template, 6 slides).

Usage: venv/bin/python deck/build_deck.py <out.pptx>
Data lives in this file; slide 4 takes deck/assets/minisite.png when present, slide 6 takes deck/option3.json when present.
"""
import copy
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
REPO = next(q for q in Path(__file__).resolve().parents if (q / ".claude/references/softserve-deck-base.pptx").exists())
BASE = str(REPO / ".claude/references/softserve-deck-base.pptx")

# brand tokens (template values, softserve-deck-kit.md)
ORANGE = RGBColor(0xF3, 0x69, 0x49)
ORANGE_T = RGBColor(0xFD, 0xED, 0xE8)
BLUE = RGBColor(0x14, 0x85, 0xC3)
BLUE_T = RGBColor(0xD2, 0xE7, 0xF6)
INK = RGBColor(0x26, 0x28, 0x2B)
GREY = RGBColor(0x80, 0x80, 0x80)
GREY_D = RGBColor(0x59, 0x59, 0x59)
HAIR = RGBColor(0xDC, 0xDC, 0xDC)
PANEL = RGBColor(0xF7, 0xF9, 0xFA)
OCI_T = RGBColor(0xEE, 0xF1, 0xF3)
OCI = RGBColor(0x6B, 0x76, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MONO = "Roboto Mono"

X0, X1 = 358732, 11852232          # content left / right edge
W = X1 - X0
TOP = 1700000                       # content top under the title
BOTTOM = 6380000                    # content bottom, clear of the bottom hairline (6515100)
MINISITE = "https://claude.ai/artifact/HTEJADBQF3ZevFPuSoTHri"
RUNNING = "Oracle packs · packaging strategy"


# ---------------------------------------------------------------- primitives
def _set_font(run, size=None, bold=None, color=None, font=None, italic=None):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    if font is not None:
        f.name = font


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, font=None, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=None, inset=0):
    """paras: list of paragraphs; a paragraph is a str or a list of (text, {fmt}) runs."""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(inset))
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing is not None and i > 0:
            p.space_before = Pt(spacing)
        if line is not None:
            p.line_spacing = line
        runs = [(para, {})] if isinstance(para, str) else para
        for t, fmt in runs:
            r = p.add_run()
            r.text = t
            _set_font(r, size=fmt.get("size", size), bold=fmt.get("bold", bold), color=fmt.get("color", color),
                      font=fmt.get("font", font), italic=fmt.get("italic"))
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def pill(slide, x, y, label, fill=None, outline=None, color=INK, size=8, h=210000, pad=90000, bold=True, w=None, font=MONO):
    est = w or int(len(label) * size * 0.62 * 12700) + 2 * pad
    s = rect(slide, x, y, est, h, fill=fill, line=outline, line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = s.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    _set_font(r, size=size, bold=bold, color=color, font=font)
    return est


def hline(slide, x, y, w, color=HAIR, width=0.75):
    c = slide.shapes.add_connector(1, Emu(x), Emu(y), Emu(x + w), Emu(y))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


# ---------------------------------------------------------------- icons (native shapes, rule 28)
def icon_badge(slide, x, y, d, kind, fill=BLUE, fg=WHITE):
    """A filled circle of diameter d with a white pictogram inside."""
    rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    c = d / 2
    if kind == "funnel":
        rect(slide, x + d * .25, y + d * .26, d * .5, d * .5, fill=fg, shape=MSO_SHAPE.FUNNEL)
    elif kind == "bolt":
        rect(slide, x + d * .30, y + d * .2, d * .4, d * .6, fill=fg, shape=MSO_SHAPE.LIGHTNING_BOLT)
    elif kind == "gear":
        rect(slide, x + d * .24, y + d * .24, d * .52, d * .52, fill=fg, shape=MSO_SHAPE.GEAR_6)
        rect(slide, x + c - d * .09, y + c - d * .09, d * .18, d * .18, fill=fill, shape=MSO_SHAPE.OVAL)
    elif kind == "target":
        for k, col in ((.56, fg), (.38, fill), (.2, fg)):
            rect(slide, x + c - d * k / 2, y + c - d * k / 2, d * k, d * k, fill=col, shape=MSO_SHAPE.OVAL)
    elif kind == "compass":
        rect(slide, x + d * .2, y + d * .2, d * .6, d * .6, fill=None, line=fg, line_w=1.5, shape=MSO_SHAPE.OVAL)
        s = rect(slide, x + c - d * .09, y + d * .27, d * .18, d * .46, fill=fg, shape=MSO_SHAPE.DIAMOND)
        s.rotation = 40
    elif kind == "sun":
        rect(slide, x + d * .22, y + d * .22, d * .56, d * .56, fill=fg, shape=MSO_SHAPE.SUN)
    elif kind == "snow":
        rect(slide, x + d * .24, y + d * .24, d * .52, d * .52, fill=fg, shape=MSO_SHAPE.STAR_6_POINT)
    elif kind == "box":
        rect(slide, x + d * .26, y + d * .26, d * .48, d * .48, fill=fg, shape=MSO_SHAPE.CUBE)


# ---------------------------------------------------------------- slide scaffolding
def strip_title_caps(run):
    rpr = run._r.get_or_add_rPr()
    rpr.set("cap", "none")
    for old in rpr.findall(qn("a:latin")):
        rpr.remove(old)
    latin = etree.SubElement(rpr, qn("a:latin"))
    latin.set("typeface", "+mj-lt")


def new_slide(prs, sldnum_xml, title, section, notes=None, title_parts=None):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            tf = ph.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            for t, accent in (title_parts or [(title, False)]):
                r = p.add_run()
                r.text = t
                r.font.size = Pt(28)
                if accent:
                    r.font.color.rgb = ORANGE
                strip_title_caps(r)
            ph.left, ph.top, ph.width, ph.height = Emu(X0), Emu(1081987), Emu(W), Emu(470000)
        elif idx == 34:
            ph.text_frame.text = f"{RUNNING} · {section}"
        else:
            ph._element.getparent().remove(ph._element)
    # slide number: copy the base's placeholder with a fresh id (else PowerPoint "repairs" the file)
    el = copy.deepcopy(sldnum_xml)
    el.find(".//" + qn("p:cNvPr")).set("id", str(s.shapes._next_shape_id))
    s.shapes._spTree.append(el)
    if notes:
        set_notes(s, notes)
    return s


NOTES_SP = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
    '<p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')


def set_notes(slide, notes):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:
        ns.shapes._spTree.append(etree.fromstring(NOTES_SP))
    ns.notes_text_frame.text = notes


# ---------------------------------------------------------------- slide 1 · goal
def slide_goal(prs, sn):
    s = new_slide(prs, sn, None, "Goal", title_parts=[("Goal: packs that ", False), ("feed Oracle's go-to-market", True)],
                  notes="The test for every pack we build or keep: does it fit how Oracle already sells, give an Oracle seller "
                        "a reason to open a meeting, and shorten the way from interest to a signed engagement?")
    pillars = [
        ("compass", ORANGE, ["Fit Oracle's", "go-to-market"],
         "Each pack lands in a campaign Oracle already runs, for customers Oracle already has."),
        ("funnel", BLUE, ["Generate", "leads"],
         "A clear story and a live demo that get an Oracle seller into a first meeting."),
        ("bolt", BLUE, ["Accelerate", "conversions"],
         "A fixed-scope proof of value on the customer's own data, then production with SoftServe."),
    ]
    gap = 300000
    cw = (W - 2 * gap) // 3
    y, h = 1950000, 3250000
    for i, (ic, col, head, body) in enumerate(pillars):
        x = X0 + i * (cw + gap)
        rect(s, x, y, cw, h, fill=PANEL)
        rect(s, x, y, cw, 60000, fill=col)
        icon_badge(s, x + 330000, y + 380000, 760000, ic, fill=col)
        text(s, x + 330000, y + 1300000, cw - 660000, 760000, head, size=22, bold=True, line=0.95)
        text(s, x + 330000, y + 2150000, cw - 660000, 1050000, [body], size=15, color=INK, line=1.08)
    # the logic between the pillars: frame -> two outcomes
    text(s, X0, 5450000, W, 400000,
         [[("Fit", {"bold": True, "color": ORANGE}), (" is the entry ticket; ", {}),
           ("more leads", {"bold": True, "color": BLUE}), (" and ", {}), ("faster conversions", {"bold": True, "color": BLUE}),
           (" are what Oracle gets from us.", {})]], size=16, color=INK)
    return s


# ---------------------------------------------------------------- slide 2 · Oracle GTM
def gtm_card(s, x, y, w, h, head, sub, badge, rows, style):
    """style: 'p1' (solid orange header), 'p2' (solid blue header), 'cold' (outlined grey)."""
    hh = 640000
    if style == "cold":
        rect(s, x, y, w, h, fill=None, line=GREY, line_w=1.0)
        rect(s, x, y, w, hh, fill=OCI_T)
        hcol, scol = INK, GREY_D
    else:
        rect(s, x, y, w, h, fill=PANEL)
        rect(s, x, y, w, hh, fill=ORANGE if style == "p1" else BLUE)
        hcol, scol = WHITE, WHITE
    text(s, x + 200000, y + 95000, w - 1500000, 300000, [head], size=15, bold=True, color=hcol)
    text(s, x + 200000, y + 370000, w - 400000, 240000, [sub], size=10, color=scol, font=MONO)
    if badge:
        bw = pill(s, 0, 0, badge, fill=WHITE, color=ORANGE if style == "p1" else BLUE, size=8, h=230000) if False else None
        lab = badge
        est = int(len(lab) * 8 * 0.62 * 12700) + 180000
        pill(s, x + w - est - 180000, y + 120000, lab, fill=WHITE, color=(ORANGE if style == "p1" else BLUE if style == "p2" else GREY_D),
             size=8, h=230000, w=est)
    ry = y + hh + 170000
    rh = (h - hh - 260000) // len(rows)
    icol = ORANGE if style == "p1" else BLUE if style == "p2" else GREY
    for k, (ic, label, body) in enumerate(rows):
        yy = ry + k * rh
        icon_badge(s, x + 200000, yy + 20000, 380000, ic, fill=icol)
        text(s, x + 700000, yy, w - 900000, 230000, [label.upper()], size=8.5, color=GREY, font=MONO)
        text(s, x + 700000, yy + 220000, w - 900000, rh - 240000, [body], size=13, color=INK, line=1.05)
        if k < len(rows) - 1:
            hline(s, x + 700000, yy + rh - 70000, w - 900000)


def slide_gtm(prs, sn):
    s = new_slide(prs, sn, None, "Oracle's go-to-market",
                  title_parts=[("Oracle brings us into ", False), ("two warm campaigns", True), ("; cold intros we develop", False)],
                  notes="Oracle AI & Data, EMEA. Warm: Oracle's laser campaigns (13 EMEA campaigns; AIDP sells to the line of business, "
                        "AI Lakehouse to technical buyers) and the NVIDIA-funded pipeline (AI-Q, cuOpt and VSS PoVs; Oracle routes the "
                        "top opportunities to SoftServe through the joint pipeline call). Cold: one-off introductions by Oracle account "
                        "teams (NATO via Oracle's defence lead, Toyota). Sources: laser-campaign deck (17 Sep), Neil + Gero sync (16 Sep), "
                        "Oscar's AIDP sessions (29 Sep, 2 and 12 Oct).")
    # zone labels
    warm_w = int(W * 0.655)
    cold_x = X0 + warm_w + 260000
    cold_w = X1 - cold_x
    y0 = 1720000
    icon_badge(s, X0, y0, 300000, "sun", fill=INK)
    text(s, X0 + 380000, y0 + 15000, warm_w, 300000,
         [[("WARM CAMPAIGNS", {"bold": True, "color": INK}), ("  ·  where Oracle AI & Data (EMEA) sees us explicitly", {"color": GREY_D})]],
         size=11, font=MONO)
    icon_badge(s, cold_x, y0, 300000, "snow", fill=GREY)
    text(s, cold_x + 380000, y0 + 15000, cold_w - 380000, 300000,
         [[("COLD", {"bold": True, "color": INK}), ("  ·  we explore proactively", {"color": GREY_D})]], size=11, font=MONO)
    cy, ch = 2150000, 4150000
    gap = 220000
    cw = (warm_w - gap) // 2
    gtm_card(s, X0, cy, cw, ch, "AIDP + AI Lakehouse", "ORACLE AI DATA PLATFORM", "PRIORITY 1", [
        ("funnel", "Campaign", "Oracle's laser campaigns"),
        ("gear", "Use cases", "Lightweight “ask your data” AI"),
        ("target", "Customers", "Oracle EBS (older ERP) estates; firms with several in-house sources on Oracle databases"),
    ], "p1")
    gtm_card(s, X0 + cw + gap, cy, cw, ch, "OCI + NVIDIA NeMo", "OCI DEDICATED AI CLUSTERS", "PRIORITY 2", [
        ("funnel", "Campaign", "NVIDIA-funded pipeline"),
        ("gear", "Use cases", "Optimisation (cuOpt), video (VSS), research agents (AI-Q)"),
        ("target", "Customers", "Large existing OCI tenants"),
    ], "p2")
    gtm_card(s, cold_x, cy, cold_w, ch, "AIDP or OCI", "AI DATA PLATFORM OR DEDICATED AI CLUSTERS", None, [
        ("funnel", "Campaign", "One-off Oracle intros to named accounts (NATO, Toyota)"),
        ("gear", "Use cases", "To be found per account"),
        ("target", "Customers", "Existing Oracle customers: EBS, Fusion, OCI"),
    ], "cold")
    return s


# ---------------------------------------------------------------- slide 3 · strategies
def slide_strategies(prs, sn):
    s = new_slide(prs, sn, None, "Packaging strategies",
                  title_parts=[("Three ways to package; ", False), ("Oracle backs the first", True)],
                  notes="Oracle's view comes from the 16 Sep sync with Neil and Gero (adopt, don't build; no AIDP customer packages yet; "
                        "keep NVIDIA work apart from AIDP) and later calls. Oscar (AIDP product strategy, 29 Sep) asked SoftServe to repeat "
                        "its NVIDIA-pack role on AIDP, so option 2 may reopen once the first AIDP opportunity is worked jointly.")
    opts = [
        ("1", "Package what we sell", "One illustrative pack per platform, plus a pack for every case we take through presales and PoVs.",
         "CURRENT", True, ("Preferred", " by Neil and Gero"), "pos", "Slide 4"),
        ("2", "Package likely bestsellers", "Proactively package popular cases that showcase what Oracle pushes.",
         "POTENTIAL", False, ("Not now", " for AIDP; little resonance with Neil and Gero"), "neg", "Slide 5"),
        ("3", "Package our own expertise", "Proactively package cases beyond databases and ERP, such as physical AI, built on our expertise.",
         "POTENTIAL", False, ("Off focus:", " mostly OCI, while Neil and Gero push AIDP and the coming Oracle Intelligence platform"), "neu", "Slide 6"),
    ]
    gap = 300000
    cw = (W - 2 * gap) // 3
    y, h = 1800000, 4500000
    for i, (num, name, what, status, current, view, tone, ref) in enumerate(opts):
        x = X0 + i * (cw + gap)
        rect(s, x, y, cw, h, fill=PANEL if current else None, line=None if current else HAIR, line_w=1.0)
        rect(s, x, y, cw, 60000, fill=ORANGE if current else GREY)
        # number badge + status
        b = rect(s, x + 280000, y + 300000, 520000, 520000, fill=None, line=INK, line_w=1.25, shape=MSO_SHAPE.OVAL)
        text(s, x + 280000, y + 300000, 520000, 520000, [num], size=20, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        est = int(len(status) * 8 * 0.62 * 12700) + 180000
        pill(s, x + cw - est - 280000, y + 430000, status, fill=ORANGE if current else None, outline=None if current else GREY_D,
             color=INK, size=8, h=240000, w=est)
        text(s, x + 280000, y + 1020000, cw - 560000, 400000, [name], size=17, bold=True)
        text(s, x + 280000, y + 1500000, cw - 560000, 1250000, [what], size=15, color=INK, line=1.08)
        hline(s, x + 280000, y + 2850000, cw - 560000)
        text(s, x + 280000, y + 2980000, cw - 560000, 230000, ["ORACLE'S VIEW"], size=8.5, color=GREY, font=MONO)
        mark = {"pos": ("●", BLUE), "neg": ("○", GREY_D), "neu": ("◐", GREY_D)}[tone]
        text(s, x + 280000, y + 3230000, cw - 560000, 900000,
             [[(mark[0] + "  ", {"color": mark[1], "bold": True}), (view[0], {"bold": True}), (view[1], {})]], size=14, line=1.05)
        text(s, x + 280000, y + h - 380000, cw - 560000, 260000, [f"→ {ref}"], size=9, color=GREY_D, font=MONO)
    return s


# ---------------------------------------------------------------- slide 4 · option 1
def slide_option1(prs, sn):
    s = new_slide(prs, sn, None, "Option 1",
                  title_parts=[("Option 1 today: ", False), ("our presales and PoV cases, packaged", True)],
                  notes="The mini-site is the single catalogue Oracle and SoftServe sellers link to: each product has its problem, "
                        "outcome, use cases by industry, technology and delivery packages, and most have an interactive walkthrough.")
    img = HERE / "assets" / "minisite.png"
    fy, fh = 1700000, 4250000
    if img.exists():
        from PIL import Image
        iw, ih = Image.open(img).size
        maxw, maxh = W, fh
        sc = min(maxw / iw, maxh / ih)
        pw, ph = int(iw * sc), int(ih * sc)
        px = X0 + (W - pw) // 2
        rect(s, px - 12700, fy - 12700, pw + 25400, ph + 25400, fill=None, line=HAIR, line_w=1.0)
        s.shapes.add_picture(str(img), Emu(px), Emu(fy), Emu(pw), Emu(ph))
    else:
        rect(s, X0, fy, W, fh, fill=PANEL, line=HAIR)
        text(s, X0, fy, W, fh, ["Mini-site screenshot (from the engagement deck) goes here"], size=14, color=GREY,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    tb = text(s, X0, 6060000, W, 300000,
              [[("Open the live mini-site products page →", {"color": BLUE, "bold": True}), ("   " + MINISITE + "#/products", {"color": GREY, "size": 9})]], size=11)
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = MINISITE + "#/products"
    tb.text_frame.paragraphs[0].runs[1].hyperlink.address = MINISITE + "#/products"
    return s


# ---------------------------------------------------------------- tables (slides 5, 6)
RUNS_STYLE = {
    "AIDP + AI Lakehouse": (ORANGE, INK),
    "AIDP": (ORANGE_T, INK),
    "OCI dedicated AI clusters": (BLUE_T, BLUE),
}


def table_slide(s, spec, rows, y=1660000, body=9.5, key=None, muted=()):
    """spec: list of (header, width, kind); kind in status | pack | text | runs | small. rows: tuples, one cell per column."""
    n = len(rows)
    hh = 260000
    bottom = BOTTOM - (230000 if key else 0)
    rh = (bottom - y - hh) // n
    xs = [X0]
    for _, w_, _ in spec[:-1]:
        xs.append(xs[-1] + w_)
    for (lab, w_, kind), x in zip(spec, xs):
        text(s, x + 50000, y, w_ - 100000, hh - 40000, [lab.upper()], size=8, color=GREY, font=MONO, anchor=MSO_ANCHOR.BOTTOM)
    hline(s, X0, y + hh, W, color=ORANGE, width=1.0)
    for k, r in enumerate(rows):
        yy = y + hh + k * rh
        for (lab, w_, kind), x, cell in zip(spec, xs, r):
            if kind == "status" and cell:
                fill, outline, col = {"NEW": (INK, None, WHITE), "HAVE": (None, GREY, GREY_D),
                                      "EXTEND": (None, INK, INK), "SERVICE": (INK, None, WHITE)}[cell]
                pill(s, x + 50000, yy + 75000, cell, fill=fill, outline=outline, color=col, size=6.5, h=180000, w=w_ - 100000)
            elif kind == "pack":
                text(s, x + 50000, yy + 55000, w_ - 100000, rh - 80000, [[(ln, {"bold": True})] for ln in cell.split("\n")],
                     size=body + 0.5, line=0.95, color=GREY_D if k in muted else INK)
            elif kind == "runs":
                fill, col = RUNS_STYLE[cell["platform"]]
                lab_ = cell["platform"]
                est = min(w_ - 100000, int(len(lab_) * 8 * 0.5 * 12700) + 170000)
                pill(s, x + 50000, yy + 65000, lab_, fill=fill, color=col, size=8, h=200000, w=est, font=None)
                if cell.get("note"):
                    text(s, x + 50000, yy + 285000, w_ - 100000, rh - 290000, [cell["note"]], size=7.5, color=GREY_D)
            elif kind in ("text", "small"):
                paras = cell if isinstance(cell, list) else [cell]
                text(s, x + 50000, yy + 55000, w_ - 100000, rh - 80000, paras, size=(body - 1 if kind == "small" else body),
                     color=(GREY_D if (kind == "small" or k in muted) else INK), line=0.98)
        hline(s, X0, yy + rh, W)
    if key:
        text(s, X0, bottom + 60000, W, 200000, key, size=8, color=GREY_D)
    return s


OPTION2 = [
    ("HAVE", "Business metrics Q&A +\nCross-system ERP Q&A",
     [[("Managers ", {"bold": True}), ("wait a week for answers across finance and sales; numbers disagree.", {})]],
     "Metrics defined once: ask in plain words, one number in minutes.",
     {"platform": "AIDP + AI Lakehouse"}, "★ Laser L1; L3, A3"),
    ("NEW", "Document Q&A",
     [[("Legal ", {"bold": True}), ("reads every contract to answer one question.", {})]],
     "Ask stored documents in plain words; the answer quotes the clause.",
     {"platform": "AIDP"}, "Laser A5"),
    ("NEW", "P&L variance analysis",
     [[("The CFO ", {"bold": True}), ("needs two weeks to explain a budget miss.", {})]],
     "Oracle's agent ranks the causes and proposes an action, same day.",
     {"platform": "AIDP"}, "Laser A2, A4"),
    ("NEW", "EBS exception agents",
     [[("EBS finance and supply ", {"bold": True}), ("find held invoices, late orders, slow stock too late.", {})]],
     "Agents flag each case and prepare the fix for approval: faster close.",
     {"platform": "AIDP"}, "★ Laser A1"),
    ("NEW", "Customer churn and retention",
     [[("Telcos and utilities ", {"bold": True}), ("spot leavers at cancellation; discounts go to stayers.", {})]],
     "Weekly risk score per customer, with the reason and the cheapest offer.",
     {"platform": "AIDP"}, "Upsell in L1, A3"),
    ("HAVE", "Plan vs actual investigation",
     [[("Project controls ", {"bold": True}), ("can't trace overruns, so the next bid repeats them.", {})]],
     "Each overrun traced to its cause and record before the next bid.",
     {"platform": "OCI dedicated AI clusters"}, "NVIDIA-funded"),
    ("EXTEND", "Large docs processing and review",
     [[("Operations ", {"bold": True}), ("retype contract rates by hand; wrong rates get paid.", {})]],
     "Terms pulled with page links; a person confirms.",
     {"platform": "OCI dedicated AI clusters", "note": "+ an AIDP edition to add"}, "NVIDIA-funded; Laser A5"),
    ("NEW", "Customer service agent",
     [[("Contact centres: ", {"bold": True}), ("scripted bots fail; routine requests queue for staff.", {})]],
     "An agent answers, completes simple requests, hands off the rest.",
     {"platform": "OCI dedicated AI clusters", "note": "with Oracle Digital Assistant"}, "★ Laser O1"),
    ("HAVE", "Workforce optimization",
     [[("Dispatchers ", {"bold": True}), ("plan by hand: uneven loads, fewer paid jobs a day.", {})]],
     "Four-week plan in minutes; more paid jobs a day.",
     {"platform": "OCI dedicated AI clusters"}, "NVIDIA-funded; Laser A3"),
    ("HAVE", "Repair-or-replace decisions",
     [[("Claims and repair teams ", {"bold": True}), ("judge damage by eye: needless replacements.", {})]],
     "Damage measured from photos against the rules; a person confirms.",
     {"platform": "OCI dedicated AI clusters"}, "NVIDIA-funded"),
    ("SERVICE", "AI running-cost reduction",
     [[("CTOs ", {"bold": True}), ("pay an outside AI service per request; data leaves the company.", {})]],
     "Apps move to a tuned open model on OCI: same answers, lower cost.",
     {"platform": "OCI dedicated AI clusters"}, "★ Laser O3"),
]


def slide_option2(prs, sn):
    s = new_slide(prs, sn, None, "Option 2",
                  title_parts=[("Option 2: new packs that ", False), ("show what Oracle's platforms do", True)],
                  notes="Built from the 13 Oct one-pager (Oracle pack set, executive version). ★ = one of Oracle's four priority laser "
                        "campaigns. Not released yet: Oracle's Deep Insights agent (whitelisted tenancies) and the EBS agent templates "
                        "(partner release ~16 Oct, per Oscar), so P&L variance analysis and EBS exception agents are built per deal until "
                        "then. Customer service agent: chat runs on Oracle Digital Assistant + OCI Generative AI; the voice tier on NVIDIA "
                        "speech models is still a prototype. Also kept, without their own row: Account insights and Case evidence collection "
                        "(same engine as Plan vs actual), Fleet route optimization (same engine as Workforce optimization).")
    spec = [("", 620000, "status"), ("Pack", 2300000, "pack"), ("Who has the problem", 3050000, "text"),
            ("What changes", 2900000, "text"), ("Runs on", 1720000, "runs"),
            ("Campaign", W - 620000 - 2300000 - 3050000 - 2900000 - 1720000, "small")]
    order = [r for r in OPTION2 if r[0] in ("NEW", "EXTEND", "SERVICE")] + [r for r in OPTION2 if r[0] == "HAVE"]
    muted = tuple(i for i, r in enumerate(order) if r[0] == "HAVE")
    key = [[("★ ", {"bold": True}), ("Oracle priority campaign   ·   ", {}), ("Laser L / A / O: ", {"bold": True}),
            ("Oracle's campaigns for AI Lakehouse, AIDP and OCI AI customers   ·   ", {}),
            ("NVIDIA-funded: ", {"bold": True}), ("OCI deals Oracle brings with NVIDIA funding   ·   ", {}),
            ("Grey rows: ", {"bold": True}), ("packs we already have", {})]]
    table_slide(s, spec, order, body=9.5, key=key, muted=muted)
    return s


def slide_option3(prs, sn):
    s = new_slide(prs, sn, None, "Option 3",
                  title_parts=[("Option 3: ", False), ("our expertise, beyond databases and ERP", True)],
                  notes="Cases from the 'SoftServe OCI/AIDP expertise use cases' session.")
    data_file = HERE / "option3.json"
    spec = [("Case", 2600000, "pack"), ("Problem", 3600000, "text"), ("Solution", 3540000, "text"),
            ("Runs on", W - 2600000 - 3600000 - 3540000, "runs")]
    if data_file.exists():
        data = json.loads(data_file.read_text())
        rows = [(c["name"], c["problem"], c["solution"], {"platform": c["platform"], "note": c.get("note")})
                for c in data["cases"]]
        if data.get("notes"):
            set_notes(s, data["notes"])
        table_slide(s, spec, rows, body=10.5)
    else:
        hh = 260000
        for (lab, w_, kind), x in zip(spec, [X0, X0 + spec[0][1], X0 + spec[0][1] + spec[1][1], X0 + spec[0][1] + spec[1][1] + spec[2][1]]):
            text(s, x + 50000, 1660000, w_ - 100000, hh - 40000, [lab.upper()], size=8, color=GREY, font=MONO, anchor=MSO_ANCHOR.BOTTOM)
        hline(s, X0, 1660000 + hh, W, color=ORANGE, width=1.0)
        rect(s, X0, 1660000 + hh + 120000, W, 1500000, fill=PANEL, line=HAIR)
        text(s, X0, 1660000 + hh + 120000, W, 1500000,
             [[("To be filled: ", {"bold": True}), ("the cases from the “SoftServe OCI/AIDP expertise use cases” session", {})],
              [("Each row: case name · problem (role, situation, cost) · solution (what changes, the measure) · "
                "runs on OCI dedicated AI clusters, AIDP or AIDP + AI Lakehouse", {"color": GREY_D, "size": 11})]],
             size=14, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=6)
    return s


def main():
    out = Path(sys.argv[1])
    prs = Presentation(BASE)
    base_slide = prs.slides[0]
    sldnum = None
    for sh in base_slide.placeholders:
        if sh.placeholder_format.idx == 4:
            sldnum = copy.deepcopy(sh._element)
    for build in (slide_goal, slide_gtm, slide_strategies, slide_option1, slide_option2, slide_option3):
        build(prs, sldnum)
    # drop the base's sample slide
    sldIdLst = prs.slides._sldIdLst
    first = sldIdLst[0]
    prs.part.drop_rel(first.rId)
    sldIdLst.remove(first)
    prs.save(out)
    print("saved", out, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
