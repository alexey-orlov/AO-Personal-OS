#!/usr/bin/env python3
"""polish_feature_list.py <in.docx> [<out.docx>] — owner's corrections of 2026-09-22 on the feature list,
applied after build until the builder carries them: (1) no generated-on footer row; (2) status glyphs read as
one size — ● and ○ are drawn smaller than ◐ in Replica/Arial, so they get a larger point size."""
import sys, docx
from docx.shared import Pt
src = sys.argv[1]; dst = sys.argv[2] if len(sys.argv) > 2 else src
d = docx.Document(src)
GL = {"●": 13.5, "○": 13.5, "◐": 10.5}        # table cells
GL_LEGEND = {"●": 10.0, "○": 10.0, "◐": 8.0}   # inline in the legend line
removed = 0
for sec in d.sections:                                    # (1) footer in the page footer
    for p in sec.footer.paragraphs:
        if "generated" in p.text or "spec v" in p.text:
            for r in p.runs: r.text = ""
            removed += 1
for p in list(d.paragraphs):                              # (1) or as a trailing body paragraph
    if "generated 20" in p.text and "feature list" in p.text.lower():
        p._element.getparent().remove(p._element); removed += 1
fixed = 0
for t in d.tables:                                        # (2) table glyphs
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    g = r.text.strip()
                    if g in GL: r.font.size = Pt(GL[g]); fixed += 1
for p in d.paragraphs:                                    # (2) legend glyphs
    for r in p.runs:
        g = r.text.strip()
        if g in GL_LEGEND and len(r.text) <= 2: r.font.size = Pt(GL_LEGEND[g]); fixed += 1
d.save(dst); print(f"polish: footer rows removed {removed}, glyph runs resized {fixed}")
